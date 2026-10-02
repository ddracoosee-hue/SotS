"""P04 ingest tests (T04.010-T04.016, 05 Stage 1)."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument

from sots.agents.failsafes.f20_invariants import InvariantContext, check_all
from sots.ingest.ingest import ingest_file, ingest_paste
from sots.ingest.loader import SCANNED_WARNING, extract_text
from sots.ingest.normalize import normalize_text
from sots.ingest.supplements import index_supplements
from sots.storage import db as storage_db
from sots.storage import repo as storage_repo

CONTENT = "Hello world under the midday sun today, bright and clear for us all here."


def _make_pdf(pages: list[str]) -> bytes:
    """Minimal valid multi-page PDF with a correct xref table."""
    objects: list[bytes] = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{4 + i * 2} 0 R" for i in range(len(pages)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for text in pages:
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET".encode()
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 3 0 R >> >> "
            f"/MediaBox [0 0 612 792] /Contents {len(objects) + 2} 0 R >>".encode()
        )
        objects.append(
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF".encode()
    )
    return bytes(out)


@pytest.fixture
def conn(tmp_path: Path):
    handle = storage_db.connect(tmp_path / "t.db")
    storage_db.migrate(handle)
    yield handle
    handle.close()


def test_same_content_same_text(tmp_path: Path) -> None:
    """T04.010: txt/md/docx/pdf with the same content extract identically."""
    assert len(CONTENT) > 50  # comfortably above the scanned-page threshold
    txt = tmp_path / "a.txt"
    txt.write_text(CONTENT, encoding="utf-8")
    md = tmp_path / "b.md"
    md.write_text(CONTENT, encoding="utf-8")
    docx_path = tmp_path / "c.docx"
    document = DocxDocument()
    document.add_paragraph(CONTENT)
    document.save(docx_path)
    pdf = tmp_path / "d.pdf"
    pdf.write_bytes(_make_pdf([CONTENT]))
    texts = [extract_text(path) for path in (txt, md, docx_path, pdf)]
    assert [item.text for item in texts] == [CONTENT] * 4
    assert all(item.warnings == [] for item in texts)


def test_docx_joins_paragraphs(tmp_path: Path) -> None:
    document = DocxDocument()
    document.add_paragraph("First.")
    document.add_paragraph("Second with caf\u00e9.")
    path = tmp_path / "c.docx"
    document.save(path)
    assert extract_text(path).text == "First.\n\nSecond with caf\u00e9."


def test_cp1252_fallback(tmp_path: Path) -> None:
    path = tmp_path / "cp.txt"
    path.write_bytes(b"caf\xe9")  # invalid UTF-8, valid cp1252
    assert extract_text(path).text == "caf\u00e9"


def test_scanned_pdf_warning(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    path.write_bytes(_make_pdf(["tiny"]))
    extracted = extract_text(path)
    assert extracted.warnings == [SCANNED_WARNING]


def test_unsupported_format(tmp_path: Path) -> None:
    path = tmp_path / "x.png"
    path.write_bytes(b"\x89PNG")
    with pytest.raises(ValueError, match="unsupported ingest format"):
        extract_text(path)


def test_normalize_endings_nbsp_blanks() -> None:
    """T04.011: endings/NBSP/blank runs normalized; words (and typos) kept."""
    raw = "one\r\ntwo\ry with\u00a0nbsp\n\n\n\n\nfive teh end"
    assert normalize_text(raw) == "one\ntwo\ny with nbsp\n\n\nfive teh end"
    assert normalize_text(raw).split() == raw.split()


def test_ingest_roundtrip_and_duplicates(tmp_path: Path, conn) -> None:
    """T04.012: inbox copy byte-identical; duplicates reuse by default."""
    inbox = tmp_path / "inbox"
    source = tmp_path / "note.md"
    source.write_bytes(b"Hello\r\n\r\n\r\n\r\nworld")
    first = ingest_file(source, conn=conn, inbox_dir=inbox, chapter_id="ch03")
    assert first.duplicate_of is None and first.warnings == []
    doc = first.document
    assert doc.chapter_id == "ch03" and doc.title == "note"
    assert (inbox / f"{doc.id}.md").read_bytes() == source.read_bytes()
    canonical = (inbox / f"{doc.id}.txt").read_text(encoding="utf-8")
    assert canonical == "Hello\n\n\nworld"
    assert doc.char_count == len(canonical) and doc.word_count == 2
    assert storage_repo.get_document(conn, doc.id) is not None

    second = ingest_file(source, conn=conn, inbox_dir=inbox)
    assert second.duplicate_of == doc.id
    assert second.document.id == doc.id

    forced = ingest_file(source, conn=conn, inbox_dir=inbox, allow_duplicate=True)
    assert forced.duplicate_of is None
    assert forced.document.id != doc.id
    assert forced.document.sha256 == doc.sha256


def test_ingest_paste(tmp_path: Path, conn) -> None:
    """T04.013: pasted text stages as paste_<ts>.md, then ingests as usual."""
    inbox = tmp_path / "inbox"
    outcome = ingest_paste("pasted\r\nwords", conn=conn, inbox_dir=inbox)
    staged = sorted(inbox.glob("paste_*.md"))
    assert len(staged) == 1
    assert outcome.document.title == staged[0].stem
    canonical = (inbox / f"{outcome.document.id}.txt").read_text(encoding="utf-8")
    assert canonical == "pasted\nwords"


def test_supplements_index(tmp_path: Path, conn) -> None:
    """T04.015: supplements indexed with text+hash; reruns are incremental."""
    supp = tmp_path / "supplements"
    (supp / "sub").mkdir(parents=True)
    (supp / "a.md").write_text("alpha", encoding="utf-8")
    (supp / "sub" / "b.txt").write_text("beta", encoding="utf-8")
    (supp / "c.png").write_bytes(b"\x89PNG")
    (supp / ".hidden.md").write_text("nope", encoding="utf-8")
    first = index_supplements(supp, conn)
    assert (first.indexed, first.unchanged) == (2, 0)
    assert first.skipped == {"c.png": "unsupported format"}
    stored = storage_repo.get_supplement_doc(conn, "a.md")
    assert stored is not None and stored.text == "alpha"
    assert len(stored.sha256) == 64

    rerun = index_supplements(supp, conn)
    assert (rerun.indexed, rerun.unchanged) == (0, 2)

    (supp / "a.md").write_text("alpha v2", encoding="utf-8")
    changed = index_supplements(supp, conn)
    assert (changed.indexed, changed.unchanged) == (1, 1)
    assert storage_repo.get_supplement_doc(conn, "a.md").text == "alpha v2"

    missing = index_supplements(tmp_path / "nope", conn)
    assert (missing.indexed, missing.unchanged) == (0, 0)


def test_raw_files_unchanged_invariant(tmp_path: Path, conn) -> None:
    """T04.016: a modified or missing inbox file is a violation."""
    inbox = tmp_path / "inbox"
    source = tmp_path / "note.md"
    source.write_text("hello", encoding="utf-8")
    outcome = ingest_file(source, conn=conn, inbox_dir=inbox)
    ctx = InvariantContext(root=tmp_path, manifest={}, conn=conn)
    assert check_all(ctx)["raw_files_unchanged"] == []

    raw_copy = Path(outcome.document.inbox_path)
    raw_copy.write_bytes(b"tampered")
    assert any(
        "sha256 mismatch" in violation
        for violation in check_all(ctx)["raw_files_unchanged"]
    )
    raw_copy.unlink()
    assert any(
        "missing" in violation for violation in check_all(ctx)["raw_files_unchanged"]
    )
